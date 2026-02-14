# src/api/Service/ExportService.py
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import os
import tempfile
from datetime import datetime
from src.utils.general.logs import HandleLogs

class ExportService:

    # ============================================
    # EXPORTACION A PDF
    # ============================================

    @staticmethod
    def export_to_pdf(data, metadata, formato='portrait'):
        """
        Exportar datos a PDF
        """
        try:
            # Crear archivo temporal
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
            temp_filename = temp_file.name
            temp_file.close()

            # Configurar documento PDF
            if formato == 'landscape':
                pagesize = (A4[1], A4[0])  # Landscape
            else:
                pagesize = A4

            doc = SimpleDocTemplate(
                temp_filename,
                pagesize=pagesize,
                rightMargin=72,
                leftMargin=72,
                topMargin=72,
                bottomMargin=18
            )

            # Estilos
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=16,
                textColor=colors.HexColor('#2c3e50'),
                alignment=TA_CENTER,
                spaceAfter=20
            )

            subtitle_style = ParagraphStyle(
                'CustomSubtitle',
                parent=styles['Heading2'],
                fontSize=12,
                textColor=colors.HexColor('#34495e'),
                alignment=TA_LEFT,
                spaceAfter=10
            )

            # Construir contenido
            story = []

            # Titulo del reporte
            title = ExportService._get_report_title(metadata.get('tipo_reporte', 'Reporte'))
            story.append(Paragraph(title, title_style))
            story.append(Spacer(1, 12))

            # Informacion del reporte - mas compacta con fechas formateadas
            info_data = []

            # Fecha de generacion formateada
            fecha_gen = metadata.get('fecha_generacion', '')
            if fecha_gen:
                try:
                    if 'T' in fecha_gen:
                        dt = datetime.fromisoformat(fecha_gen.replace('Z', '+00:00'))
                        fecha_formateada = dt.strftime('%d/%m/%Y %H:%M')
                    else:
                        fecha_formateada = fecha_gen
                    info_data.append(['Generado:', f"{fecha_formateada} por {metadata.get('generado_por', 'Sistema')}"])
                except Exception:
                    info_data.append(['Generado por:', metadata.get('generado_por', 'Sistema')])
            else:
                info_data.append(['Generado por:', metadata.get('generado_por', 'Sistema')])

            info_data.append(['Total Registros:', str(metadata.get('total_registros', 0))])

            # Agregar filtros aplicados de forma mas clara
            if metadata.get('filtros_aplicados'):
                filtros = metadata['filtros_aplicados']
                if filtros.get('id_paciente'):
                    info_data.append(['Filtro:', f"Paciente ID {filtros['id_paciente']}"])
                elif filtros.get('id_personal'):
                    info_data.append(['Filtro:', f"Personal ID {filtros['id_personal']}"])

            info_table = Table(info_data, colWidths=[1.5*inch, 2.5*inch])
            info_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f8f9fa')),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ]))

            story.append(info_table)
            story.append(Spacer(1, 15))

            # Datos del reporte
            if data and len(data) > 0:
                story.append(Paragraph("Datos del Reporte", subtitle_style))

                # Convertir datos a tabla
                table_data = ExportService._prepare_table_data_for_pdf(data, metadata.get('tipo_reporte'))

                if table_data and len(table_data) > 1:  # Al menos headers + 1 fila de datos
                    num_cols = len(table_data[0])
                    available_width = pagesize[0] - 144  # Margins (72*2)

                    # Calcular anchos de columna segun el tipo de reporte
                    if metadata.get('tipo_reporte') == 'asistencia_paciente':
                        # Anchos especificos para asistencia de paciente (9 columnas)
                        total_weight = sum([20, 12, 20, 15, 18, 8, 8, 8, 9])
                        col_widths = [
                            available_width * (20/total_weight),  # Paciente
                            available_width * (12/total_weight),  # Cedula
                            available_width * (20/total_weight),  # Sesion
                            available_width * (15/total_weight),  # Especialidad
                            available_width * (18/total_weight),  # Terapeuta
                            available_width * (8/total_weight),   # Sesiones
                            available_width * (8/total_weight),   # Asistidas
                            available_width * (8/total_weight),   # Perdidas
                            available_width * (9/total_weight)    # % Asist.
                        ]
                    elif metadata.get('tipo_reporte') == 'progreso_terapeutico':
                        if num_cols <= 7:
                            weights = [20, 18, 15, 18, 12, 8, 9][:num_cols]
                        else:
                            weights = [15, 12, 12, 12, 10, 8, 8, 8, 8, 8, 8, 8][:num_cols]
                        total_weight = sum(weights)
                        col_widths = [available_width * (w/total_weight) for w in weights]
                    elif metadata.get('tipo_reporte') == 'academico_estudiante':
                        # Anchos para reporte academico (9 columnas en PDF)
                        weights = [18, 12, 18, 15, 15, 8, 8, 8, 9][:num_cols]
                        total_weight = sum(weights)
                        col_widths = [available_width * (w/total_weight) for w in weights]
                    elif metadata.get('tipo_reporte') == 'carga_trabajo_personal':
                        weights = [20, 15, 10, 10, 12, 10][:num_cols]
                        total_weight = sum(weights)
                        col_widths = [available_width * (w/total_weight) for w in weights]
                    elif metadata.get('tipo_reporte') == 'rendimiento_clase':
                        weights = [18, 15, 15, 10, 10, 10, 10, 10][:num_cols]
                        total_weight = sum(weights)
                        col_widths = [available_width * (w/total_weight) for w in weights]
                    elif metadata.get('tipo_reporte') == 'utilizacion_recursos':
                        weights = [20, 15, 12, 12, 12, 10, 10][:num_cols]
                        total_weight = sum(weights)
                        col_widths = [available_width * (w/total_weight) for w in weights]
                    elif metadata.get('tipo_reporte') == 'estadisticas_generales':
                        weights = [25, 15, 15, 15, 15, 15][:num_cols]
                        total_weight = sum(weights)
                        col_widths = [available_width * (w/total_weight) for w in weights]
                    else:
                        # Distribucion para reportes sin mapping especifico
                        if num_cols <= 6:
                            col_width = available_width / num_cols
                            col_widths = [col_width] * num_cols
                        else:
                            base_width = available_width / num_cols
                            col_widths = []
                            for i in range(num_cols):
                                if i % 2 == 0:
                                    col_widths.append(base_width * 1.3)
                                else:
                                    col_widths.append(base_width * 0.7)

                    data_table = Table(table_data, colWidths=col_widths)
                    data_table.setStyle(TableStyle([
                        # Header styling
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, 0), 9),
                        # Data rows styling
                        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                        ('FONTSIZE', (0, 1), (-1, -1), 8),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                        ('LEFTPADDING', (0, 0), (-1, -1), 3),
                        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
                        ('TOPPADDING', (0, 0), (-1, -1), 4),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                        # Alternating row colors
                        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
                        # Alineacion especifica para columnas numericas
                        ('ALIGN', (-3, 1), (-1, -1), 'CENTER'),
                    ]))

                    story.append(data_table)
                else:
                    story.append(Paragraph("Los datos no se pudieron procesar correctamente.", styles['Normal']))
            else:
                story.append(Paragraph("No se encontraron datos para el reporte", styles['Normal']))

            # Generar PDF
            doc.build(story)

            return {
                'success': True,
                'file_path': temp_filename,
                'filename': f"reporte_{metadata.get('tipo_reporte', 'general')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            }

        except Exception as e:
            HandleLogs.write_error(f"Error en export_to_pdf: {str(e)}")
            # Limpiar archivo temporal en caso de error
            if 'temp_filename' in locals():
                ExportService.cleanup_temp_file(temp_filename)
            return {
                'success': False,
                'message': f'Error generando PDF: {str(e)}'
            }

    # ============================================
    # EXPORTACION A EXCEL
    # ============================================

    @staticmethod
    def export_to_excel(data, metadata):
        """
        Exportar datos a Excel
        """
        try:
            # Crear archivo temporal
            temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
            temp_filename = temp_file.name
            temp_file.close()

            # Crear workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Reporte"

            # Estilos
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="3498DB", end_color="3498DB", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")

            info_font = Font(bold=True)
            info_fill = PatternFill(start_color="ECF0F1", end_color="ECF0F1", fill_type="solid")

            # Titulo del reporte
            title = ExportService._get_report_title(metadata.get('tipo_reporte', 'Reporte'))
            ws['A1'] = title
            ws['A1'].font = Font(bold=True, size=16)
            ws.merge_cells('A1:F1')

            # Informacion del reporte
            current_row = 3
            info_items = [
                ('Fecha de Generacion:', metadata.get('fecha_generacion', 'N/A')),
                ('Generado por:', metadata.get('generado_por', 'Sistema')),
                ('Total de Registros:', str(metadata.get('total_registros', 0)))
            ]

            if metadata.get('filtros_aplicados'):
                filtros = metadata['filtros_aplicados']
                if filtros.get('fecha_inicio'):
                    info_items.append(('Fecha Inicio:', str(filtros['fecha_inicio'])))
                if filtros.get('fecha_fin'):
                    info_items.append(('Fecha Fin:', str(filtros['fecha_fin'])))

            for label, value in info_items:
                ws[f'A{current_row}'] = label
                ws[f'B{current_row}'] = value
                ws[f'A{current_row}'].font = info_font
                ws[f'A{current_row}'].fill = info_fill
                current_row += 1

            # Datos del reporte
            if data and len(data) > 0:
                current_row += 2
                ws[f'A{current_row}'] = "Datos del Reporte"
                ws[f'A{current_row}'].font = Font(bold=True, size=14)
                current_row += 1

                # Headers y datos
                headers, rows = ExportService._prepare_table_data_for_excel(data, metadata.get('tipo_reporte'))

                # Escribir headers
                for col_index, header in enumerate(headers, 1):
                    cell = ws.cell(row=current_row, column=col_index, value=header)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment

                current_row += 1

                # Escribir datos
                for row_data in rows:
                    for col_index, value in enumerate(row_data, 1):
                        ws.cell(row=current_row, column=col_index, value=value)
                    current_row += 1

                # Ajustar ancho de columnas
                for col_num in range(1, len(headers) + 1):
                    max_length = 0
                    column_letter = get_column_letter(col_num)

                    # Revisar todas las celdas de la columna que contienen datos
                    for row in range(current_row - len(rows), current_row):
                        try:
                            cell = ws.cell(row=row, column=col_num)
                            if hasattr(cell, 'value') and cell.value is not None:
                                cell_length = len(str(cell.value))
                                if cell_length > max_length:
                                    max_length = cell_length
                        except Exception:
                            pass

                    # Tambien revisar el header
                    try:
                        header_length = len(str(headers[col_num - 1]))
                        if header_length > max_length:
                            max_length = header_length
                    except Exception:
                        pass

                    adjusted_width = min(max(max_length + 2, 10), 50)
                    ws.column_dimensions[column_letter].width = adjusted_width

            # Guardar archivo
            wb.save(temp_filename)

            return {
                'success': True,
                'file_path': temp_filename,
                'filename': f"reporte_{metadata.get('tipo_reporte', 'general')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            }

        except Exception as e:
            HandleLogs.write_error(f"Error en export_to_excel: {str(e)}")
            # Limpiar archivo temporal en caso de error
            if 'temp_filename' in locals():
                ExportService.cleanup_temp_file(temp_filename)
            return {
                'success': False,
                'message': f'Error generando Excel: {str(e)}'
            }

    # ============================================
    # METODOS AUXILIARES PRIVADOS
    # ============================================

    @staticmethod
    def _get_report_title(tipo_reporte):
        """Obtener titulo legible del reporte"""
        titles = {
            'asistencia_paciente': 'Reporte de Asistencia por Paciente',
            'progreso_terapeutico': 'Reporte de Progreso Terapeutico',
            'carga_trabajo_personal': 'Reporte de Carga de Trabajo del Personal',
            'academico_estudiante': 'Reporte Academico por Estudiante',
            'rendimiento_clase': 'Reporte de Rendimiento por Clase',
            'utilizacion_recursos': 'Reporte de Utilizacion de Recursos',
            'estadisticas_generales': 'Estadisticas Generales del Centro'
        }
        return titles.get(tipo_reporte, 'Reporte del Sistema')

    @staticmethod
    def _format_date_value(value):
        """Formatear valor de fecha para exportacion"""
        try:
            if value and 'GMT' in str(value):
                dt = datetime.strptime(str(value), '%a, %d %b %Y %H:%M:%S GMT')
                return dt.strftime('%d/%m/%Y')
            elif value:
                return str(value)[:10]
            else:
                return ''
        except Exception:
            return str(value)[:10] if value else ''

    @staticmethod
    def _prepare_table_data_for_pdf(data, tipo_reporte):
        """Preparar datos para tabla PDF"""
        if not data:
            return []

        # Headers y mapeo especifico por tipo de reporte
        if tipo_reporte == 'asistencia_paciente':
            headers = ['Paciente', 'Cedula', 'Sesion', 'Especialidad', 'Terapeuta', 'Sesiones', 'Asistidas', 'Perdidas', '% Asist.']
            field_mapping = [
                'paciente_nombre',
                'cedula',
                'sesion_titulo',
                'especialidad',
                'terapeuta_nombre',
                'total_sesiones_programadas',
                'sesiones_asistidas',
                'sesiones_perdidas',
                'porcentaje_asistencia'
            ]
        elif tipo_reporte == 'progreso_terapeutico':
            headers = ['Paciente', 'Sesion', 'Especialidad', 'Estado', 'Ingreso', 'Primera', 'Ultima']
            field_mapping = [
                'paciente_nombre', 'sesion_titulo', 'especialidad', 'estado_tratamiento',
                'fecha_ingreso', 'fecha_primera_sesion', 'fecha_ultima_sesion'
            ]
        elif tipo_reporte == 'carga_trabajo_personal':
            headers = ['Terapeuta', 'Especialidad', 'S.Activas', 'Pacientes', 'Programadas', 'Realizadas']
            field_mapping = [
                'terapeuta_nombre', 'especialidad', 'sesiones_activas', 'total_pacientes',
                'sesiones_programadas', 'sesiones_realizadas'
            ]
        elif tipo_reporte == 'academico_estudiante':
            headers = ['Estudiante', 'Cedula', 'Clase', 'Especialidad', 'Educador', 'Clases', 'Asistidas', 'Perdidas', '% Asist.']
            field_mapping = [
                'estudiante_nombre', 'cedula', 'clase_titulo', 'especialidad',
                'educador_nombre', 'total_clases_programadas', 'clases_asistidas',
                'clases_perdidas', 'porcentaje_asistencia'
            ]
        elif tipo_reporte == 'rendimiento_clase':
            headers = ['Clase', 'Especialidad', 'Educador', 'Estudiantes', 'Prom. Asist.', 'Prom. Calif.', 'Aprobados', 'Reprobados']
            field_mapping = [
                'clase_titulo', 'especialidad', 'educador_nombre', 'total_estudiantes',
                'promedio_asistencia', 'promedio_calificacion', 'aprobados', 'reprobados'
            ]
        elif tipo_reporte == 'utilizacion_recursos':
            headers = ['Recurso', 'Tipo', 'Capacidad', 'Uso Actual', '% Utilizacion', 'Disponible', 'Estado']
            field_mapping = [
                'recurso_nombre', 'tipo_recurso', 'capacidad', 'uso_actual',
                'porcentaje_utilizacion', 'disponible', 'estado'
            ]
        elif tipo_reporte == 'estadisticas_generales':
            headers = ['Indicador', 'Valor', 'Periodo', 'Variacion', 'Estado']
            field_mapping = [
                'indicador', 'valor', 'periodo', 'variacion', 'estado'
            ]
        else:
            # Para otros tipos de reporte, crear mapeo dinamico
            if data and len(data) > 0:
                first_row = data[0]
                if isinstance(first_row, dict):
                    field_mapping = list(first_row.keys())[:8]  # Maximo 8 columnas para PDF
                    headers = [key.replace('_', ' ').title() for key in field_mapping]
                else:
                    headers = ['Campo', 'Valor']
                    field_mapping = []
            else:
                headers = ['Campo', 'Valor']
                field_mapping = []

        # Campos que son fechas (para formateo automatico)
        date_fields = {
            'fecha_ingreso', 'fecha_primera_sesion', 'fecha_ultima_sesion',
            'fecha_inicio', 'fecha_fin', 'fecha_creacion', 'fecha_modificacion'
        }

        # Campos que son porcentajes
        percentage_fields = {
            'porcentaje_asistencia', 'promedio_asistencia', 'porcentaje_utilizacion'
        }

        # Campos de texto largo (truncar)
        long_text_fields = {
            'paciente_nombre': 20, 'estudiante_nombre': 20, 'sesion_titulo': 20,
            'clase_titulo': 20, 'objetivo_general': 20, 'recurso_nombre': 20,
            'indicador': 20, 'especialidad': 15, 'terapeuta_nombre': 15,
            'educador_nombre': 15, 'cedula': 12, 'estado_tratamiento': 10
        }

        # Preparar filas de datos
        table_data = [headers]

        for row in data:
            if isinstance(row, dict):
                row_data = []
                for field in field_mapping:
                    value = row.get(field, '')

                    if field in percentage_fields:
                        try:
                            row_data.append(f"{float(value):.1f}%")
                        except (ValueError, TypeError):
                            row_data.append('0.0%')
                    elif field in date_fields:
                        row_data.append(ExportService._format_date_value(value))
                    elif field in long_text_fields:
                        max_len = long_text_fields[field]
                        row_data.append(str(value)[:max_len] if value else '')
                    else:
                        row_data.append(str(value) if value else '')

                table_data.append(row_data)
            else:
                # Fallback para datos en formato tuple/list
                row_data = []
                for i, value in enumerate(row):
                    if i < len(headers):
                        if isinstance(value, (int, float)):
                            row_data.append(str(value))
                        else:
                            row_data.append(str(value)[:20] if value else '')
                    else:
                        break
                table_data.append(row_data)

        return table_data

    @staticmethod
    def _prepare_table_data_for_excel(data, tipo_reporte):
        """Preparar datos para Excel usando el mismo mapeo que PDF"""
        if not data:
            return [], []

        # Usar el mismo sistema de mapeo que PDF pero con headers mas descriptivos para Excel
        if tipo_reporte == 'asistencia_paciente':
            headers = [
                'Paciente', 'ID Paciente', 'Cedula', 'Sesion', 'Especialidad', 'Terapeuta',
                'Total Sesiones Programadas', 'Sesiones Asistidas', 'Sesiones Perdidas',
                'Sesiones Canceladas', 'Porcentaje Asistencia'
            ]
            field_mapping = [
                'paciente_nombre', 'paciente_id', 'cedula', 'sesion_titulo',
                'especialidad', 'terapeuta_nombre', 'total_sesiones_programadas',
                'sesiones_asistidas', 'sesiones_perdidas', 'sesiones_canceladas',
                'porcentaje_asistencia'
            ]
        elif tipo_reporte == 'progreso_terapeutico':
            headers = [
                'Paciente', 'ID Paciente', 'Sesion', 'Especialidad', 'Estado Tratamiento',
                'Fecha Ingreso', 'Fecha Primera Sesion', 'Fecha Ultima Sesion', 'Objetivo General'
            ]
            field_mapping = [
                'paciente_nombre', 'paciente_id', 'sesion_titulo', 'especialidad',
                'estado_tratamiento', 'fecha_ingreso', 'fecha_primera_sesion',
                'fecha_ultima_sesion', 'objetivo_general'
            ]
        elif tipo_reporte == 'carga_trabajo_personal':
            headers = [
                'Terapeuta', 'ID Personal', 'Especialidad', 'Total Sesiones Activas',
                'Total Pacientes', 'Sesiones Programadas', 'Sesiones Realizadas',
                'Sesiones Canceladas', 'Horas Promedio Sesion', 'Total Horas Trabajadas'
            ]
            field_mapping = [
                'terapeuta_nombre', 'personal_id', 'especialidad', 'sesiones_activas',
                'total_pacientes', 'sesiones_programadas', 'sesiones_realizadas',
                'sesiones_canceladas', 'horas_promedio_sesion', 'total_horas_trabajadas'
            ]
        elif tipo_reporte == 'academico_estudiante':
            headers = [
                'Estudiante', 'ID Estudiante', 'Cedula', 'Clase', 'Especialidad', 'Educador',
                'Total Clases Programadas', 'Clases Asistidas', 'Clases Perdidas',
                'Porcentaje Asistencia', 'Calificacion Promedio', 'Evaluaciones Aprobadas'
            ]
            field_mapping = [
                'estudiante_nombre', 'estudiante_id', 'cedula', 'clase_titulo',
                'especialidad', 'educador_nombre', 'total_clases_programadas',
                'clases_asistidas', 'clases_perdidas', 'porcentaje_asistencia',
                'calificacion_promedio', 'evaluaciones_aprobadas'
            ]
        elif tipo_reporte == 'rendimiento_clase':
            headers = [
                'Clase', 'ID Clase', 'Especialidad', 'Educador', 'Total Estudiantes',
                'Promedio Asistencia', 'Promedio Calificacion', 'Aprobados', 'Reprobados'
            ]
            field_mapping = [
                'clase_titulo', 'clase_id', 'especialidad', 'educador_nombre',
                'total_estudiantes', 'promedio_asistencia', 'promedio_calificacion',
                'aprobados', 'reprobados'
            ]
        elif tipo_reporte == 'utilizacion_recursos':
            headers = [
                'Recurso', 'Tipo de Recurso', 'Capacidad', 'Uso Actual',
                'Porcentaje Utilizacion', 'Disponible', 'Estado'
            ]
            field_mapping = [
                'recurso_nombre', 'tipo_recurso', 'capacidad', 'uso_actual',
                'porcentaje_utilizacion', 'disponible', 'estado'
            ]
        elif tipo_reporte == 'estadisticas_generales':
            headers = [
                'Indicador', 'Valor', 'Periodo', 'Variacion', 'Estado'
            ]
            field_mapping = [
                'indicador', 'valor', 'periodo', 'variacion', 'estado'
            ]
        else:
            # Para otros tipos de reporte, crear mapeo dinamico
            if data and len(data) > 0:
                first_row = data[0]
                if isinstance(first_row, dict):
                    field_mapping = list(first_row.keys())
                    headers = [key.replace('_', ' ').title() for key in field_mapping]
                else:
                    headers = ['Campo ' + str(i+1) for i in range(len(first_row))]
                    field_mapping = list(range(len(first_row)))
            else:
                headers = ['Sin Datos']
                field_mapping = []

        # Campos que son fechas y porcentajes
        date_fields = {
            'fecha_ingreso', 'fecha_primera_sesion', 'fecha_ultima_sesion',
            'fecha_inicio', 'fecha_fin', 'fecha_creacion', 'fecha_modificacion'
        }
        percentage_fields = {
            'porcentaje_asistencia', 'promedio_asistencia', 'porcentaje_utilizacion',
            'calificacion_promedio', 'promedio_calificacion'
        }

        # Preparar filas de datos usando el mapeo de campos
        rows = []
        for row in data:
            if isinstance(row, dict):
                row_data = []
                for field in field_mapping:
                    value = row.get(field, '')

                    if field in percentage_fields:
                        try:
                            row_data.append(float(value))
                        except (ValueError, TypeError):
                            row_data.append(0.0)
                    elif field in date_fields:
                        row_data.append(ExportService._format_date_value(value))
                    else:
                        row_data.append(value)

                rows.append(row_data)
            else:
                # Fallback para datos en formato tuple/list
                rows.append(list(row))

        return headers, rows

    @staticmethod
    def cleanup_temp_file(file_path):
        """Limpiar archivo temporal"""
        try:
            if file_path and os.path.exists(file_path):
                os.unlink(file_path)
                return True
        except Exception as e:
            HandleLogs.write_error(f"Error eliminando archivo temporal {file_path}: {str(e)}")
        return False
